"""
Report generation service with templates, scheduling, and export
"""
import logging
from datetime import datetime, timezone, timedelta
from typing import Dict, List, Optional
from enum import Enum
import json

logger = logging.getLogger(__name__)

class ReportFormat(Enum):
    """Report export formats"""
    PDF = 'pdf'
    EXCEL = 'excel'
    CSV = 'csv'
    JSON = 'json'

class ReportTemplate:
    """Report template definition"""
    
    def __init__(self, name: str, description: str, sections: List[str]):
        self.name = name
        self.description = description
        self.sections = sections  # e.g., ['summary', 'users', 'parameters', 'alerts']
        self.created_at = datetime.now(timezone.utc)
    
    def to_dict(self) -> dict:
        return {
            'name': self.name,
            'description': self.description,
            'sections': self.sections,
            'created_at': self.created_at.isoformat()
        }

class Report:
    """Generated report"""
    
    def __init__(self, report_id: int, template_name: str, data: dict, created_by: str):
        self.report_id = report_id
        self.template_name = template_name
        self.data = data
        self.created_by = created_by
        self.created_at = datetime.now(timezone.utc)
        self.shared_with = []
    
    def share_with(self, user_email: str):
        """Share report with user"""
        if user_email not in self.shared_with:
            self.shared_with.append(user_email)
            logger.info(f"[REPORT] Report {self.report_id} shared with {user_email}")
    
    def to_dict(self) -> dict:
        return {
            'report_id': self.report_id,
            'template_name': self.template_name,
            'data': self.data,
            'created_by': self.created_by,
            'created_at': self.created_at.isoformat(),
            'shared_with': self.shared_with
        }

class ScheduledReport:
    """Scheduled report generation"""
    
    def __init__(self, schedule_id: int, template_name: str, frequency: str, created_by: str):
        self.schedule_id = schedule_id
        self.template_name = template_name
        self.frequency = frequency  # 'daily', 'weekly', 'monthly'
        self.created_by = created_by
        self.created_at = datetime.now(timezone.utc)
        self.last_generated = None
        self.next_generation = self._calculate_next_generation()
        self.enabled = True
    
    def _calculate_next_generation(self) -> datetime:
        """Calculate next generation time"""
        now = datetime.now(timezone.utc)
        if self.frequency == 'daily':
            return now + timedelta(days=1)
        elif self.frequency == 'weekly':
            return now + timedelta(weeks=1)
        elif self.frequency == 'monthly':
            return now + timedelta(days=30)
        return now
    
    def should_generate(self) -> bool:
        """Check if report should be generated"""
        return self.enabled and datetime.now(timezone.utc) >= self.next_generation
    
    def mark_generated(self):
        """Mark as generated and update next generation time"""
        self.last_generated = datetime.now(timezone.utc)
        self.next_generation = self._calculate_next_generation()
    
    def to_dict(self) -> dict:
        return {
            'schedule_id': self.schedule_id,
            'template_name': self.template_name,
            'frequency': self.frequency,
            'created_by': self.created_by,
            'created_at': self.created_at.isoformat(),
            'last_generated': self.last_generated.isoformat() if self.last_generated else None,
            'next_generation': self.next_generation.isoformat(),
            'enabled': self.enabled
        }

class ReportGenerationService:
    """Generates and manages reports"""
    
    def __init__(self):
        self.templates: Dict[str, ReportTemplate] = {}
        self.reports: Dict[int, Report] = {}
        self.scheduled_reports: Dict[int, ScheduledReport] = {}
        self.report_counter = 0
        self.schedule_counter = 0
        self._init_default_templates()
    
    def _init_default_templates(self):
        """Initialize default report templates"""
        templates = [
            ReportTemplate('System Summary', 'Overview of system status', ['summary', 'users', 'parameters']),
            ReportTemplate('User Activity', 'User creation, updates, and deletions', ['users', 'audit']),
            ReportTemplate('Parameter Report', 'Parameter values and history', ['parameters', 'alerts']),
            ReportTemplate('Audit Report', 'Complete audit log', ['audit']),
            ReportTemplate('Full Report', 'Complete system report', ['summary', 'users', 'parameters', 'alerts', 'audit'])
        ]
        
        for template in templates:
            self.templates[template.name] = template
    
    def create_template(self, name: str, description: str, sections: List[str]) -> tuple[bool, Optional[str]]:
        """Create custom report template"""
        try:
            if name in self.templates:
                return False, "Template already exists"
            
            template = ReportTemplate(name, description, sections)
            self.templates[name] = template
            
            logger.info(f"[REPORT] Created template {name}")
            return True, None
        except Exception as e:
            logger.error(f"[REPORT] Error creating template: {e}")
            return False, str(e)
    
    def get_template(self, name: str) -> Optional[ReportTemplate]:
        """Get report template"""
        return self.templates.get(name)
    
    def get_all_templates(self) -> List[dict]:
        """Get all templates"""
        return [t.to_dict() for t in self.templates.values()]
    
    def generate_report(self, template_name: str, data: dict, created_by: str) -> tuple[bool, Optional[int], Optional[str]]:
        """Generate report from template"""
        try:
            if template_name not in self.templates:
                return False, None, "Template not found"
            
            self.report_counter += 1
            report = Report(self.report_counter, template_name, data, created_by)
            self.reports[self.report_counter] = report
            
            logger.info(f"[REPORT] Generated report {self.report_counter} from template {template_name}")
            return True, self.report_counter, None
        except Exception as e:
            logger.error(f"[REPORT] Error generating report: {e}")
            return False, None, str(e)
    
    def get_report(self, report_id: int) -> Optional[Report]:
        """Get report by ID"""
        return self.reports.get(report_id)
    
    def get_all_reports(self) -> List[dict]:
        """Get all reports"""
        return [r.to_dict() for r in self.reports.values()]
    
    def share_report(self, report_id: int, user_email: str) -> tuple[bool, Optional[str]]:
        """Share report with user"""
        try:
            if report_id not in self.reports:
                return False, "Report not found"
            
            self.reports[report_id].share_with(user_email)
            return True, None
        except Exception as e:
            logger.error(f"[REPORT] Error sharing report: {e}")
            return False, str(e)
    
    def export_report(self, report_id: int, format: ReportFormat) -> tuple[bool, Optional[bytes], Optional[str]]:
        """Export report to specified format"""
        try:
            if report_id not in self.reports:
                return False, None, "Report not found"
            
            report = self.reports[report_id]
            
            if format == ReportFormat.JSON:
                data = json.dumps(report.to_dict(), indent=2).encode()
                return True, data, None
            
            elif format == ReportFormat.CSV:
                # Simple CSV export
                import csv
                import io
                output = io.StringIO()
                writer = csv.writer(output)
                
                # Write header
                writer.writerow(['Key', 'Value'])
                
                # Write data
                for key, value in report.data.items():
                    writer.writerow([key, json.dumps(value) if isinstance(value, (dict, list)) else value])
                
                return True, output.getvalue().encode(), None
            
            elif format == ReportFormat.EXCEL:
                # Excel export requires openpyxl
                try:
                    import openpyxl
                    from openpyxl.utils import get_column_letter
                    
                    wb = openpyxl.Workbook()
                    ws = wb.active
                    ws.title = "Report"
                    
                    # Write header
                    ws['A1'] = 'Key'
                    ws['B1'] = 'Value'
                    
                    # Write data
                    row = 2
                    for key, value in report.data.items():
                        ws[f'A{row}'] = key
                        ws[f'B{row}'] = json.dumps(value) if isinstance(value, (dict, list)) else value
                        row += 1
                    
                    # Save to bytes
                    import io
                    output = io.BytesIO()
                    wb.save(output)
                    output.seek(0)
                    return True, output.getvalue(), None
                except ImportError:
                    return False, None, "openpyxl not installed"
            
            elif format == ReportFormat.PDF:
                # PDF export requires reportlab
                try:
                    from reportlab.lib.pagesizes import letter
                    from reportlab.pdfgen import canvas
                    import io
                    
                    output = io.BytesIO()
                    c = canvas.Canvas(output, pagesize=letter)
                    
                    # Write title
                    c.setFont("Helvetica-Bold", 16)
                    c.drawString(50, 750, f"Report: {report.template_name}")
                    
                    # Write metadata
                    c.setFont("Helvetica", 10)
                    c.drawString(50, 730, f"Generated: {report.created_at.isoformat()}")
                    c.drawString(50, 715, f"Created by: {report.created_by}")
                    
                    # Write data
                    y = 700
                    for key, value in report.data.items():
                        text = f"{key}: {json.dumps(value) if isinstance(value, (dict, list)) else value}"
                        c.drawString(50, y, text[:80])  # Truncate long lines
                        y -= 15
                    
                    c.save()
                    output.seek(0)
                    return True, output.getvalue(), None
                except ImportError:
                    return False, None, "reportlab not installed"
            
            return False, None, f"Unsupported format: {format}"
        except Exception as e:
            logger.error(f"[REPORT] Error exporting report: {e}")
            return False, None, str(e)
    
    def schedule_report(self, template_name: str, frequency: str, created_by: str) -> tuple[bool, Optional[int], Optional[str]]:
        """Schedule report generation"""
        try:
            if template_name not in self.templates:
                return False, None, "Template not found"
            
            if frequency not in ['daily', 'weekly', 'monthly']:
                return False, None, "Invalid frequency"
            
            self.schedule_counter += 1
            scheduled = ScheduledReport(self.schedule_counter, template_name, frequency, created_by)
            self.scheduled_reports[self.schedule_counter] = scheduled
            
            logger.info(f"[REPORT] Scheduled {frequency} report for template {template_name}")
            return True, self.schedule_counter, None
        except Exception as e:
            logger.error(f"[REPORT] Error scheduling report: {e}")
            return False, None, str(e)
    
    def get_scheduled_reports(self) -> List[dict]:
        """Get all scheduled reports"""
        return [r.to_dict() for r in self.scheduled_reports.values()]
    
    def process_scheduled_reports(self) -> List[int]:
        """Process due scheduled reports"""
        generated_ids = []
        try:
            for schedule_id, scheduled in self.scheduled_reports.items():
                if scheduled.should_generate():
                    # Generate report (simplified - in real implementation, fetch actual data)
                    success, report_id, error = self.generate_report(
                        scheduled.template_name,
                        {'generated_at': datetime.now(timezone.utc).isoformat()},
                        'system'
                    )
                    if success:
                        scheduled.mark_generated()
                        generated_ids.append(report_id)
                        logger.info(f"[REPORT] Auto-generated report {report_id}")
        except Exception as e:
            logger.error(f"[REPORT] Error processing scheduled reports: {e}")
        
        return generated_ids

# Global instance
_report_service = None

def get_report_generation_service() -> ReportGenerationService:
    """Get or create report generation service"""
    global _report_service
    if _report_service is None:
        _report_service = ReportGenerationService()
    return _report_service

# when MQTT disconnects mark staus indicator as "Disconnected" and streamed data from desktop stores in local_buffer table of sqlite database of ** @dspl-precision-pulse-desktop**  and mark it as unsynced in parameter_stream table of both postgres and sqlite database of @dspl-precision-pulse-backend  and ** @dspl-precision-pulse-desktop** respectively. also stops the data streming on web
# after reconnection of mqtt mark the status indicator as "Connected" and stored data from local_buffer maek it as synced and stream that synced data to web app before desktop streams new data and after streaming data flush the entire data from local_buffer table